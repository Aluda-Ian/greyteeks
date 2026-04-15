import csv
import json
from io import TextIOWrapper

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, HttpResponseNotAllowed, JsonResponse
from django.shortcuts import render, redirect
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt
from openpyxl import load_workbook

from .models import Contact, Group, LeadForm

@login_required
def contact_list(request):
    """List the current user's contacts with optional group filtering."""
    groups = Group.objects.filter(user=request.user).order_by('-created_at')
    group_id = request.GET.get('group')
    contacts = Contact.objects.filter(group__user=request.user).order_by('-created_at')
    selected_group = None

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        phone_number = request.POST.get('phone_number', '').strip()
        email = request.POST.get('email', '').strip()
        group_id = request.POST.get('group')

        if not phone_number:
            messages.error(request, 'Phone number is required to add a contact.')
        else:
            target_group = Group.objects.filter(id=group_id, user=request.user).first()
            if not target_group:
                messages.error(request, 'Please select a valid group for the contact.')
            else:
                Contact.objects.create(
                    group=target_group,
                    phone_number=phone_number,
                    name=name,
                    email=email,
                )
                messages.success(request, 'Contact added successfully.')
                return redirect('contact_list')


    if group_id:
        selected_group = groups.filter(id=group_id).first()
        if selected_group:
            contacts = contacts.filter(group=selected_group)

    return render(request, 'contacts/contact_list.html', {
        'groups': groups,
        'contacts': contacts,
        'selected_group': selected_group,
    })

@login_required
def edit_contact(request, contact_id):
    """Edit a single contact belonging to the current user."""
    contact = Contact.objects.filter(id=contact_id, group__user=request.user).first()
    if not contact:
        messages.error(request, 'Contact not found.')
        return redirect('contact_list')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        phone_number = request.POST.get('phone_number', '').strip()
        email = request.POST.get('email', '').strip()
        group_id = request.POST.get('group')

        if not phone_number:
            messages.error(request, 'Phone number is required.')
            return render(request, 'contacts/contact_edit.html', {
                'contact': contact,
                'groups': Group.objects.filter(user=request.user),
            })

        target_group = Group.objects.filter(id=group_id, user=request.user).first()
        if not target_group:
            messages.error(request, 'Please select a valid group.')
            return render(request, 'contacts/contact_edit.html', {
                'contact': contact,
                'groups': Group.objects.filter(user=request.user),
            })

        contact.name = name
        contact.phone_number = phone_number
        contact.email = email
        contact.group = target_group
        contact.save()
        messages.success(request, 'Contact updated successfully.')
        return redirect('contact_list')

    return render(request, 'contacts/contact_edit.html', {
        'contact': contact,
        'groups': Group.objects.filter(user=request.user),
    })

@login_required
def delete_contact(request, contact_id):
    """Delete a contact after confirmation."""
    contact = Contact.objects.filter(id=contact_id, group__user=request.user).first()
    if not contact:
        messages.error(request, 'Contact not found.')
        return redirect('contact_list')

    if request.method == 'POST':
        contact.delete()
        messages.success(request, 'Contact deleted successfully.')
        return redirect('contact_list')

    return render(request, 'contacts/contact_delete_confirm.html', {'contact': contact})

@login_required
def upload_contacts(request):
    """Upload a CSV or XLSX file to import contacts into a user group."""
    groups = Group.objects.filter(user=request.user).order_by('-created_at')

    if request.method == 'POST':
        group_id = request.POST.get('group')
        contact_file = request.FILES.get('contact_file')

        if not group_id or not contact_file:
            messages.error(request, "Please choose a group and upload a file.")
            return render(request, 'contacts/upload.html', {'groups': groups})

        group = Group.objects.filter(id=group_id, user=request.user).first()
        if not group:
            messages.error(request, "Selected group not found.")
            return render(request, 'contacts/upload.html', {'groups': groups})

        try:
            if contact_file.name.lower().endswith('.csv'):
                file_data = TextIOWrapper(contact_file.file, encoding='utf-8-sig')
                reader = csv.reader(file_data)
                rows = list(reader)
            elif contact_file.name.lower().endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
                workbook = load_workbook(contact_file, read_only=True, data_only=True)
                sheet = workbook.active
                rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
            else:
                raise ValueError("Unsupported file type. Please upload a .csv or .xlsx file.")
        except Exception as exc:
            messages.error(request, f"Unable to parse the file: {exc}")
            return render(request, 'contacts/upload.html', {'groups': groups})

        created_count = 0
        for row in rows:
            if not row:
                continue
            phone = str(row[0]).strip() if row[0] is not None else ''
            name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ''
            email = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ''
            if phone:
                Contact.objects.create(group=group, phone_number=phone, name=name, email=email)
                created_count += 1

        messages.success(request, f"Imported {created_count} contacts into {group.name}.")
        return redirect('dashboard')

    return render(request, 'contacts/upload.html', {'groups': groups})


@login_required
def manage_lead_forms(request):
    groups = Group.objects.filter(user=request.user).order_by('-created_at')
    lead_forms = LeadForm.objects.filter(user=request.user).order_by('-created_at')
    preview_code = None

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        target_group_id = request.POST.get('target_group')
        success_message = request.POST.get('success_message', '').strip() or 'Thank you for subscribing!'

        if not name:
            messages.error(request, 'Please give this lead form a name.')
        else:
            target_group = groups.filter(id=target_group_id).first()
            if not target_group:
                messages.error(request, 'Please choose a valid group for new leads.')
            else:
                lead_form = LeadForm.objects.create(
                    user=request.user,
                    target_group=target_group,
                    name=name,
                    success_message=success_message,
                )
                messages.success(request, 'Lead form created successfully. Copy the embed snippet below.')
                lead_forms = LeadForm.objects.filter(user=request.user).order_by('-created_at')
                preview_code = _build_lead_form_embed(request, lead_form)

    lead_forms_with_embed = [
        {
            'lead_form': lead_form,
            'embed_code': _build_lead_form_embed(request, lead_form),
            'embed_url': request.build_absolute_uri(reverse('submit_lead_form', args=[lead_form.public_uuid])),
        }
        for lead_form in lead_forms
    ]

    return render(request, 'contacts/lead_forms.html', {
        'groups': groups,
        'lead_forms': lead_forms_with_embed,
        'preview_code': preview_code,
    })


def _build_lead_form_embed(request, lead_form):
    action_url = request.build_absolute_uri(reverse('submit_lead_form', args=[lead_form.public_uuid]))
    return (
        f'<form action="{action_url}" method="POST" style="max-width:420px; font-family:system-ui, sans-serif;">\n'
        '  <div style="display:flex; flex-direction:column; gap:10px;">\n'
        '    <label style="font-weight:600;">Name\n'
        '      <input type="text" name="name" required style="width:100%; padding:10px; border:1px solid #d1d5db; border-radius:8px;">\n'
        '    </label>\n'
        '    <label style="font-weight:600;">Email\n'
        '      <input type="email" name="email" required style="width:100%; padding:10px; border:1px solid #d1d5db; border-radius:8px;">\n'
        '    </label>\n'
        '    <label style="font-weight:600;">Phone\n'
        '      <input type="tel" name="phone" style="width:100%; padding:10px; border:1px solid #d1d5db; border-radius:8px;">\n'
        '    </label>\n'
        '    <button type="submit" style="background:#3182ce; color:white; border:none; padding:12px 18px; border-radius:10px; cursor:pointer;">Subscribe</button>\n'
        '  </div>\n'
        '</form>'
    )


@csrf_exempt
def submit_lead_form(request, form_uuid):
    cors_headers = {
        'Access-Control-Allow-Origin': '*',
        'Access-Control-Allow-Methods': 'POST, OPTIONS',
        'Access-Control-Allow-Headers': 'Content-Type, X-Requested-With',
    }

    if request.method == 'OPTIONS':
        return HttpResponse(status=204, headers=cors_headers)

    if request.method != 'POST':
        response = HttpResponseNotAllowed(['POST', 'OPTIONS'])
        for header, value in cors_headers.items():
            response[header] = value
        return response

    lead_form = LeadForm.objects.filter(public_uuid=form_uuid).first()
    if not lead_form:
        return JsonResponse({'detail': 'Form not found.'}, status=404, headers=cors_headers)

    if not lead_form.is_active:
        return JsonResponse({'detail': 'This form is inactive.'}, status=403, headers=cors_headers)

    if request.content_type and request.content_type.startswith('application/json'):
        try:
            payload = json.loads(request.body.decode('utf-8') or '{}')
        except json.JSONDecodeError:
            payload = {}
        name = payload.get('name', '').strip()
        email = payload.get('email', '').strip()
        phone = payload.get('phone', '').strip()
    else:
        name = request.POST.get('name', '').strip()
        email = request.POST.get('email', '').strip()
        phone = request.POST.get('phone', '').strip()

    if not phone and not email:
        return JsonResponse({'detail': 'Email or phone is required.'}, status=400, headers=cors_headers)

    contact, created = Contact.objects.get_or_create(
        group=lead_form.target_group,
        phone_number=phone or '',
        email=email or '',
        defaults={'name': name},
    )

    if not created and name and not contact.name:
        contact.name = name
        contact.save(update_fields=['name'])

    response = JsonResponse({'success': True, 'message': lead_form.success_message}, status=200)
    for header, value in cors_headers.items():
        response[header] = value
    return response
